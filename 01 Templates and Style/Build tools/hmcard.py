"""Shared A5 card builder (extracted from build_a5_cards.py)."""
import sys, os, re, zipfile
sys.path.insert(0, os.path.dirname(__file__))
from hmdocx import Builder, DEEP, TEAL, PALE, PALETEAL, DARK, PRINCIPLES, PRINCIPLE_DEFS, RGB
from docx.shared import Pt, Twips, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

A5_W, A5_H, MARGIN = 8391, 11906, 680
BODY = A5_W - 2 * MARGIN  # 7031
SCALE = A5_W / 11906


class Card(Builder):
    def __init__(self, template, out, title):
        super().__init__(template, out, doc_title=title)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Twips(A5_W), Twips(A5_H)
        sec.left_margin = sec.right_margin = Twips(MARGIN)
        sec.top_margin, sec.bottom_margin = Twips(1000), Twips(760)
        sp = sec._sectPr
        tp = sp.find(qn("w:titlePg"))
        if tp is not None: sp.remove(tp)
        # scale the slim header and footer tables to A5
        for part in (sec.header.part, sec.footer.part):
            root = part._element
            for el in root.iter():
                for attr in (qn("w:w"),):
                    v = el.get(attr)
                    if v and el.tag in (qn("w:tblW"), qn("w:gridCol"), qn("w:tcW"), qn("w:tblInd")):
                        el.set(attr, str(int(round(int(v) * SCALE))))
                if el.tag == qn("w:tcMar"):
                    for m in el:
                        if m.get(qn("w:w")): m.set(qn("w:w"), str(MARGIN))
                if el.tag == qn("w:tab") and el.get(qn("w:pos")):
                    el.set(qn("w:pos"), str(BODY))
        # no page numbers on a printed card: drop the "Page X of Y" runs from the footer
        froot = sec.footer.part._element
        for p_ in list(froot.iter(qn("w:p"))):
            txt = "".join(t.text or "" for t in p_.iter(qn("w:t")))
            if "Page" in txt or p_.find(qn("w:fldSimple")) is not None or any(True for _ in p_.iter(qn("w:instrText"))):
                for r in list(p_):
                    if r.tag == qn("w:r") or r.tag == qn("w:fldSimple"):
                        p_.remove(r)
        # body font a touch smaller for a card
        st = self.doc.styles["Normal"]; st.font.size = Pt(8.5)
        self.doc.styles["List Bullet"].font.size = Pt(8.5)

    def side_title(self, kicker, title, new_page=False):
        p = self.para(kicker, style="DocumentType", after=0)
        if new_page: p.paragraph_format.page_break_before = True
        for r in p.runs: r.font.size = Pt(8.5)
        t = self.para(title, style="Title", after=4)
        for r in t.runs: r.font.size = Pt(18)
        return t

    def h(self, text, color=DEEP):
        p = self.para(text, bold=True, color=color, size=10, before=3, after=1, keep=True)
        return p

    def body(self, text, after=2):
        return self.para(text, size=8.5, after=after)

    def bl(self, items, size=8.5):
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            self._runs(p, it, size=size)
            p.paragraph_format.space_after = Pt(0.5)

    def box(self, lines, fill=PALE, bar=TEAL):
        t = self.doc.add_table(rows=1, cols=1)
        self._widths(t, [BODY]); self._borders(t, {"left": ("single", 18, bar)}); self._margins(t, 50, 50, 120, 120)
        c = t.cell(0, 0); self._shade(c, fill)
        first = True
        for line in lines:
            p = c.paragraphs[0] if first else c.add_paragraph(); first = False
            p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
            self._runs(p, line, color=DARK, size=8.5)
        self._cant_split(t); self.spacer(3)

    def two(self, rows, widths=None, header=False, size=8.5):
        widths = widths or [1800, BODY - 1800]
        t = self.doc.add_table(rows=len(rows), cols=2, style="Mentoring Table")
        self._widths(t, widths)
        for i, (a, b) in enumerate(rows):
            ca, cb = t.cell(i, 0), t.cell(i, 1)
            if header and i == 0:
                self.cell_text(ca, a, bold=True, color="FFFFFF", size=size); self.cell_text(cb, b, bold=True, color="FFFFFF", size=size)
                self._shade(ca, DEEP); self._shade(cb, DEEP)
            else:
                if isinstance(a, list): self.cell_bullets(ca, a, after=1)
                else: self.cell_text(ca, a, size=size)
                if isinstance(b, list): self.cell_bullets(cb, b, after=1)
                else: self.cell_text(cb, b, size=size)
            for c in (ca, cb):
                for p in c.paragraphs:
                    p.paragraph_format.space_before = Pt(0.5); p.paragraph_format.space_after = Pt(0.5)
                    for r in p.runs: r.font.size = Pt(size)
        self._look(t, 1 if header else 0, 0 if header else 1); self._cant_split(t); self.spacer(3)

    def strip(self):
        t = self.doc.add_table(rows=2, cols=4)
        self._widths(t, [BODY // 4] * 4); self._borders(t, {"insideV": ("single", 8, "FFFFFF")}); self._margins(t, 50, 50, 90, 90)
        for i, (name, fill) in enumerate(PRINCIPLES):
            c = t.cell(0, i); self._shade(c, fill)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            r = p.add_run(f"{i+1} "); r.font.color.rgb = RGB("8FD3D1"); r.bold = True; r.font.size = Pt(8.5)
            r = p.add_run(name.upper()); r.font.color.rgb = RGB("FFFFFF"); r.bold = True; r.font.size = Pt(8.5)
            c2 = t.cell(1, i); self._shade(c2, PALE)
            p2 = c2.paragraphs[0]; p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.space_before = Pt(0)
            r = p2.add_run(PRINCIPLE_DEFS[name]); r.font.size = Pt(7); r.font.color.rgb = RGB(DARK)
        self._cant_split(t); self.spacer(3)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def foot(self):
        self.para("Hertfordshire Provincial Mentoring · mentor@hertfordshirefreemasons.org.uk · solomon.ugle.org.uk · fhlbc.co.uk  |  Prepare • Guide • Support • Enlighten",
                  size=7, color=DARK, before=3, after=0)


