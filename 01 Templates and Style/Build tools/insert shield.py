import sys
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
src, out, shield = sys.argv[1], sys.argv[2], sys.argv[3]
doc = Document(src)
done = 0
for sec in doc.sections:
    for hdr in (sec.first_page_header, sec.header):
        for tbl in hdr.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    if '[Shield]' in cell.text:
                        p = cell.paragraphs[0]
                        for r in list(p.runs): r._r.getparent().remove(r._r)
                        p.add_run().add_picture(shield, height=Cm(1.3))
                        tcPr = cell._tc.tcPr
                        b = tcPr.find(qn('w:tcBorders'))
                        if b is not None: tcPr.remove(b)
                        done += 1
doc.save(out); print('inserted', done, '->', out)
