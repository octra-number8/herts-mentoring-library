"""Shared helpers for building Hertfordshire Mentoring documents on the Word template with python-docx.

Usage:
    from hmdocx import Builder
    b = Builder(template_path, out_path, doc_title="...", issue_line="Issue 1 (Draft for Review)", date_line="October 2026")
    b.para(...); b.h1("1", "..."); b.bullet(...); b.table(...); b.panel(...); b.principles_strip(); b.save()
"""
import re, copy, zipfile, os
from docx import Document
from docx.shared import Pt, Twips, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK, DEEP, DARK, TEAL, PALE, PALETEAL, MINT = "1E2329", "003668", "06345D", "157A7C", "F2F7FC", "E6F2F2", "8FD3D1"
PRINCIPLES = [("Prepare", "003668"), ("Guide", "0A4F72"), ("Support", "106576"), ("Enlighten", "157A7C")]
PRINCIPLE_DEFS = {
    "Prepare": "Ready for each stage before it happens. No surprises.",
    "Guide": "Accurate, honest, timely information. Never dependency.",
    "Support": "Present, responsive, compassionate. “How can we help?”",
    "Enlighten": "Understanding, curiosity and enjoyment, for life.",
}
STRAP_OLD = "Prepare – Support – Guide – Enlighten"
STRAP_NEW = "Prepare • Guide • Support • Enlighten"
FULL_WIDTH = 9412


def RGB(h):
    return RGBColor.from_string(h)


class Builder:
    def __init__(self, template, out, doc_title, issue_line="Issue 1 (Draft for Review)", date_line="October 2026"):
        self.out = out
        self.tmp = out + ".stage.docx"
        self._patch_package(template, doc_title, issue_line, date_line)
        self.doc = Document(self.tmp)
        self._reset_body()

    # ---------- package-level patches (header/footer text) ----------
    def _patch_package(self, template, doc_title, issue_line, date_line):
        def patch(name, text):
            if name == "word/header1.xml" and STRAP_OLD in text:
                text = text.replace(STRAP_OLD, "@@STRAP@@")
                text = text.replace(">Support<", ">@@G@@<").replace(">Guide<", ">Support<").replace(">@@G@@<", ">Guide<")
                text = text.replace("@@STRAP@@", STRAP_NEW)
            if name == "word/footer2.xml":
                text = text.replace(STRAP_OLD, STRAP_NEW)
                text = text.replace("Version [1.0]", issue_line).replace("Review date [DD/MM/YYYY]", date_line)
            if name == "word/header2.xml":
                text = text.replace("[Document title or subject]", doc_title)
            return text
        zi = zipfile.ZipFile(template); zo = zipfile.ZipFile(self.tmp, "w", zipfile.ZIP_DEFLATED)
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                data = patch(item.filename, data.decode("utf-8")).encode("utf-8")
            zo.writestr(item, data)
        zo.close(); zi.close()

    def _reset_body(self):
        body = self.doc.element.body
        layout2 = None
        for p in body.findall(qn("w:p")):
            sp = p.find(qn("w:pPr") + "/" + qn("w:sectPr"))
            if sp is not None and sp.find(qn("w:titlePg")) is not None:
                ftrs = [f.get(qn("r:id")) for f in sp.findall(qn("w:footerReference"))]
                if "rId8" in ftrs:
                    layout2 = copy.deepcopy(sp)
        assert layout2 is not None, "Layout 2 section properties not found in template"
        for f in layout2.findall(qn("w:footerReference")):
            if f.get(qn("w:type")) == "first":
                f.set(qn("r:id"), "rId7")
        for child in list(body):
            body.remove(child)
        body.append(layout2)

    # ---------- text ----------
    def _runs(self, p, text, bold=False, italic=False, color=None, size=None):
        parts = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text)
        for part in parts:
            if not part:
                continue
            b, i = bold, italic
            if part.startswith("**") and part.endswith("**"):
                part = part[2:-2]; b = True
            elif part.startswith("*") and part.endswith("*"):
                part = part[1:-1]; i = True
            r = p.add_run(part)
            r.bold = b or None
            r.italic = i or None
            if color: r.font.color.rgb = RGB(color)
            if size: r.font.size = Pt(size)
        return p

    def para(self, text="", style=None, bold=False, italic=False, color=None, size=None, after=None, before=None, align=None, keep=False):
        p = self.doc.add_paragraph(style=style)
        if text:
            self._runs(p, text, bold, italic, color, size)
        pf = p.paragraph_format
        if after is not None: pf.space_after = Pt(after)
        if before is not None: pf.space_before = Pt(before)
        if align: p.alignment = align
        if keep: pf.keep_with_next = True
        return p

    def spacer(self, pt=4):
        return self.para(after=0, size=pt)

    def h1(self, num, text, fill=DEEP):
        p = self.doc.add_paragraph(style="Heading1")
        if num:
            r = p.add_run(f" {num} ")
            rpr = r._r.get_or_add_rPr()
            c = OxmlElement("w:color"); c.set(qn("w:val"), "FFFFFF"); rpr.append(c)
            s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); rpr.append(s)
            p.add_run("  " + text)
        else:
            p.add_run(text)
        return p

    def h2(self, text):
        return self.doc.add_paragraph(text, style="Heading2")

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet"); self._runs(p, text); return p

    def bullets(self, items):
        for it in items: self.bullet(it)

    # ---------- tables ----------
    @staticmethod
    def _shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)

    @staticmethod
    def _margins(tbl, top=100, bottom=100, left=170, right=170):
        m = OxmlElement("w:tblCellMar")
        for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
        tbl._tbl.tblPr.append(m)

    @staticmethod
    def _borders(tbl, spec):
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            val, sz, col = spec.get(side, ("nil", 0, "auto"))
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), col); b.append(e)
        tbl._tbl.tblPr.append(b)

    @staticmethod
    def _widths(tbl, widths):
        tbl.autofit = False
        grid = tbl._tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(w))
        for row in tbl.rows:
            for c, w in zip(row.cells, widths): c.width = Twips(w)
        tblPr = tbl._tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
        tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)

    @staticmethod
    def _look(tbl, first_row=0, first_col=1):
        tblPr = tbl._tbl.tblPr
        tl = tblPr.find(qn("w:tblLook"))
        if tl is None:
            tl = OxmlElement("w:tblLook"); tblPr.append(tl)
        val = (0x0020 if first_row else 0) | (0x0080 if first_col else 0) | 0x0400 | 0x0200
        tl.set(qn("w:val"), "%04X" % val)
        tl.set(qn("w:firstRow"), str(first_row)); tl.set(qn("w:lastRow"), "0")
        tl.set(qn("w:firstColumn"), str(first_col)); tl.set(qn("w:lastColumn"), "0")
        tl.set(qn("w:noHBand"), "1"); tl.set(qn("w:noVBand"), "1")

    @staticmethod
    def _cant_split(tbl):
        for row in tbl.rows:
            trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))

    def cell_text(self, cell, text, bold=False, color=None, size=None, first=True):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
        self._runs(p, text, bold=bold, color=color, size=size)
        return p

    def cell_bullets(self, cell, items, after=2):
        first = True
        for it in items:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.style = self.doc.styles["List Bullet"]
            p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
            self._runs(p, it)

    def table(self, rows, widths, header=False, bullets_col=None, keep_header=True, row_height=None):
        """rows: list of lists of str (or list-of-str for bullet cells). Mentoring Table style, label column shaded."""
        t = self.doc.add_table(rows=len(rows), cols=len(widths), style="Mentoring Table")
        self._widths(t, widths)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                if header and i == 0:
                    self.cell_text(c, val, bold=True, color="FFFFFF"); self._shade(c, DEEP)
                elif isinstance(val, list):
                    self.cell_bullets(c, val)
                else:
                    self.cell_text(c, val)
        self._look(t, 1 if header else 0, 0 if header else 1)
        if row_height:
            for i, r in enumerate(t.rows):
                if header and i == 0: continue
                trPr = r._tr.get_or_add_trPr(); h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(row_height)); h.set(qn("w:hRule"), "atLeast"); trPr.append(h)
        self._cant_split(t)
        if header and keep_header:
            for c in t.rows[0].cells:
                for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
        self.spacer()
        return t

    def panel(self, lines, fill=PALE, bar=TEAL):
        t = self.doc.add_table(rows=1, cols=1)
        self._widths(t, [FULL_WIDTH])
        self._borders(t, {"left": ("single", 24, bar)})
        self._margins(t, 90, 90, 200, 200)
        c = t.cell(0, 0); self._shade(c, fill)
        first = True
        for line in lines:
            p = c.paragraphs[0] if first else c.add_paragraph(); first = False
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            self._runs(p, line, color=DARK)
        self._cant_split(t)
        self.spacer()
        return t

    def principles_strip(self):
        t = self.doc.add_table(rows=2, cols=4)
        self._widths(t, [2353] * 4)
        self._borders(t, {"insideV": ("single", 12, "FFFFFF")})
        self._margins(t, 90, 90, 140, 140)
        for i, (name, fill) in enumerate(PRINCIPLES):
            c = t.cell(0, i); self._shade(c, fill)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            r = p.add_run(f"{i+1}  "); r.font.color.rgb = RGB(MINT); r.bold = True; r.font.size = Pt(11)
            r = p.add_run(name.upper()); r.font.color.rgb = RGB("FFFFFF"); r.bold = True; r.font.size = Pt(11)
            c2 = t.cell(1, i); self._shade(c2, PALE)
            p2 = c2.paragraphs[0]; p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.space_before = Pt(0)
            r = p2.add_run(PRINCIPLE_DEFS[name]); r.font.size = Pt(9.5); r.font.color.rgb = RGB(DARK)
        self._cant_split(t)
        self.spacer()
        return t

    # ---------- standard blocks ----------
    def opening(self, doc_type, title, purpose, meta_rows):
        self.para(doc_type, style="DocumentType")
        self.para(title, style="Title")
        self.para(purpose, style="PurposeStatement")
        self.table(meta_rows, [2041, 7371])

    def closing(self, companions):
        self.para("Provincial Office, Fleet House, 10 Parkway, St Albans AL3 6PA · 01727 833061 · www.pgherts.org · solomon.ugle.org.uk  |  "
                  f"Companion to: {companions}  |  © Provincial Grand Mentor, Province of Hertfordshire, 2026",
                  size=8.5, color=DARK, before=4, after=0)

    def save(self):
        self.doc.save(self.out)
        os.remove(self.tmp)
        return self.out
