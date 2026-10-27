"""B0 The Hertfordshire Mentoring Handbook: one reference for every mentor, compiled from the library.
Usage: build_handbook.py <template.docx> <out_dir>"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from hmdocx import Builder, DEEP, DARK, TEAL, PALE, PALETEAL, MINT, FULL_WIDTH
from docx.shared import Pt
from docx.enum.text import WD_BREAK

template, out_dir = sys.argv[1], sys.argv[2]
W = FULL_WIDTH


class HB(Builder):
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def part(self, num, title, blurb):
        self.page_break()
        p = self.para(f"PART {num}", style="DocumentType", after=0)
        t = self.para(title, style="Title", after=6)
        for r in t.runs: r.font.size = Pt(30)
        self.para(blurb, style="PurposeStatement")
        self.spacer(6)

    def table(self, rows, widths, **kw):
        t = super().table(rows, widths, **kw)
        if len(rows) <= 6:  # short tables stay on one page
            for r in t.rows[:-1]:
                for c in r.cells:
                    for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
        return t

    def key(self, text, label="Key principle"):
        self.panel([f"**{label}.** {text}"], fill=PALE, bar=DEEP)

    def good(self, text):
        self.panel([f"**Good practice.** {text}"], fill=PALETEAL, bar=TEAL)

    def quote(self, text, source):
        self.panel([f"*{text}*", f"— {source}"], fill="FFFFFF", bar=MINT)


def build():
    out = os.path.join(out_dir, "B0 The Hertfordshire Mentoring Handbook Issue 1 Draft.docx")
    b = HB(template, out, doc_title="The Hertfordshire Mentoring Handbook")

    # ---------------------------------------------------------------- cover / opening
    b.opening("Practitioner handbooks · B0", "The Hertfordshire Mentoring Handbook",
              "One reference for every mentor in the Province: the principles, the people, the member’s journey from first enquiry to beyond the Chair, "
              "the moments that go wrong and what to do about them, the rules that bite, and the words and letters a new Brother will hear. "
              "Compiled from the Hertfordshire mentoring library so that a Lodge Mentor, a Personal Mentor, a Master or any Brother who has a "
              "member to look after can find what he needs in one place.",
              [("Audience", "Lodge Mentors, Personal Mentors, Worshipful Masters and Wardens, Almoners, Membership Officers, Secretaries, and any Brother supporting another"),
               ("Issue", "Issue 1 (Draft for Review), October 2026 · Produced by the Provincial Grand Mentor’s office"),
               ("Status", "Compendium. Where another document in the library owns a topic in full, this Handbook carries the working version and names the owner-document."),
               ("Standard", "Quotes the four principles verbatim from The Hertfordshire Mentoring Principles; follows the Style, Terminology and Data Guide; never reveals ritual.")])

    b.h2("Foreword")
    b.para("If you have a Brother to look after, this book is for you. It does not matter whether you were appointed Lodge Mentor at the last Installation, "
           "asked to keep an eye on the man you proposed, or find yourself in the Chair wondering what the newest member makes of your Lodge. "
           "The job is the same: to make sure that no Brother walks alone.", after=6)
    b.para("Over the past year the Province has built a library of mentoring guides, each written for one reader and one moment: a short guide for Personal "
           "Mentors, a fuller one for Lodge Mentors, checklists for the nights that matter, a card of warning signs, a guide to what happens when a member "
           "drifts or moves away, and courses to go with them. This Handbook gathers the working content of all of them into one place, in one order, "
           "so that you can read it from the front or open it at the page you need tonight.", after=6)
    b.para("Two ideas run through every page. Freemasonry must be fun: a man who is enjoying himself attends, participates, belongs, and does not resign. "
           "And Family, Work, Masonry, in that order, always: the member managing a busy life is not disengaged, and Freemasonry must fit into his life, "
           "not demand that his life fit around it.", after=6)
    b.para("Thank you for the time and care you give to this work. Hertfordshire Freemasonry is stronger because of mentors like you.", after=8)
    b.para("**Wes Pierce**", after=0)
    b.para("*Deputy Provincial Grand Mentor, Province of Hertfordshire · October 2026*", after=10)
    b.panel(["**Freemasonry must be fun.** A man who is enjoying himself attends. A man who attends participates. A man who participates belongs. And a man who belongs does not resign."], fill=PALE, bar=DEEP)

    b.h2("How to use this Handbook")
    b.table([
        ["Part One · Foundations", "What mentoring is in Hertfordshire, the four principles, who does what, and the craft of a good mentoring conversation. Read once; return when something feels off."],
        ["Part Two · The journey", "Stage by stage, from first enquiry to beyond the Chair. Open at the stage your Brother has reached. Each chapter ends with what to do this week."],
        ["Part Three · When it goes wrong", "Warning signs, the four steps, resignation and retrieval, keeping the door open, and relocation. The pages to open the day a letter arrives."],
        ["Part Four · Learning and the Province", "Supporting learning, Solomon, the Provincial support available to you, and the rhythm of a mentor’s year."],
        ["Part Five · Reference", "The rules every mentor should know, dress and etiquette answers, new members’ questions, the glossary of words and letters, contacts and the library."],
    ], [2600, 6812])
    b.h2("The library this Handbook draws on")
    b.para("Each chapter names its owner-document in the library. When a topic changes, the owner-document changes first and this Handbook follows at its next issue. "
           "Short guides and cards remain the right thing to hand to a mentor who needs one page, not a book.", after=4)
    b.table([
        ["A1", "The Hertfordshire Mentoring Principles", "The one-page canon, quoted in Chapter 2."],
        ["A2", "Provincial Mentoring Framework", "The Provincial strategy: team, New Initiate Pathway, One Journey, reporting."],
        ["A3", "Style, Terminology and Data Guide", "House style, terminology, the ritual standard, data rules for mentoring records."],
        ["B1", "Lodge Mentor’s Guide (Issue 2) and Short Guide", "The Lodge Mentor’s role and year."],
        ["B2", "Personal Mentor’s Guide and Quick Guide", "The one-to-one companion."],
        ["B6", "Retention, Retrieval and Migration Guide", "Owner-document for Part Three."],
        ["C1, C2", "Considering Freemasonry; the New Brother’s Guide", "Member-facing companions for stages 1 to 3."],
        ["D1 to D4", "Personal Mentoring Plan; Mentor Meeting Guide; Stage Checklists; Warning Signs card", "The working tools."],
        ["E1a, E2a, E6a", "Lodge Mentor Essentials; Personal Mentor Course; The Chair and Beyond", "Training, with the New Masons Evening."],
        ["E5", "Ceremony Explained series", "What each ceremony asks and means, without revealing ritual."],
        ["F1", "In That Order", "The book: the philosophy and the stories behind all of it."],
    ], [900, 3900, 4612], header=False)

    # ================================================================ PART ONE
    b.part("ONE", "Foundations", "What mentoring is in Hertfordshire, the four principles that organise it, the people who carry it, and the craft of a conversation that helps.")

    b.h1("1", "The Hertfordshire way")
    b.quote("Every member, in every Lodge, experiences Freemasonry as welcoming, supportive, engaging, and personally rewarding.", "Our vision, The Hertfordshire Mentoring Principles")
    b.para("Mentoring in Hertfordshire is a culture, not a programme. It is older than the office that now carries its name: Freemasons have always guided "
           "each other. What the Province has done is to make that guidance deliberate, consistent and available to every member, from the first enquiry "
           "to the years beyond the Chair, so that it no longer depends on the luck of who happened to sit next to him.", after=6)
    b.para("Seven propositions describe the approach. Everything else in this Handbook is one of them applied to a particular moment.", after=4)
    b.table([
        ["Enjoyment is the floor, not the ceiling", "Freemasonry must be fun. A man who enjoys his Freemasonry attends; a man who attends participates; a man who participates belongs; and a man who belongs does not resign."],
        ["Retention is an outcome, not a target", "We do not retain members by design; we build a culture that prepares and supports them, and retention follows. The question is never “How do we keep members?” but “How do we make Freemasonry so rewarding that men never want to leave?”"],
        ["Mentoring is a culture, not a programme", "The Lodge Mentor co-ordinates; he does not do it all himself. We are all mentors to others in Freemasonry. No one is left in the corner."],
        ["The member sets the pace", "No pressure: to take office, to learn ritual, to progress between degrees, or to join the Royal Arch. Information, not recruitment; timing, not inevitability; readiness, not convenience."],
        ["Family, Work, Masonry: in that order, always", "Honest expectation-setting begins before the interview and never stops. There should never be any surprises. A member managing a busy life is not disengaged."],
        ["The journey is lifelong and mentoring is circular", "Support runs from first enquiry to beyond the Chair, and flows both ways: the mentee teaches the mentor. The risk of loss does not end at the Chair."],
        ["Help, never interrogate", "The question is always “How can we help?”, never “Why aren’t you attending?”. Compassionate enquiry underpins every intervention."],
    ], [3000, 6412])
    b.key("Nobody is born a good mentor. The qualities that make one are developed, and they begin with an honest look at oneself: the rough ashlar is a metaphor for every Mason, at every stage, including the most experienced.")

    b.h1("2", "The four principles")
    b.para("The four principles below are quoted from The Hertfordshire Mentoring Principles, the Province’s canonical one-page statement. Every mentoring "
           "document in Hertfordshire quotes the same text; none restates it from memory. They are on the Mentor’s Jewel and in the Guiding Principles artwork, "
           "and they are the shape of every conversation in this book.", after=6)
    b.principles_strip()
    b.table([
        ["PREPARE", "Ensuring members are ready for each stage of their Masonic journey. This begins before interview and continues through practical preparation for ceremonies, office, and progression - ensuring members understand what to expect before it happens. There should never be any surprises."],
        ["GUIDE", "Providing accurate, clear, and timely information and support. Good guidance is honest and grounded in authoritative sources. It dispels myths, answers questions openly, and helps members navigate an unfamiliar world with confidence - without creating dependency."],
        ["SUPPORT", "Being present, responsive, and compassionate throughout. Support is not only practical; it is relational. Members who feel cared for and valued are far more likely to remain engaged and committed. The question is always “How can we help?” - never “Why aren’t you attending?”"],
        ["ENLIGHTEN", "Deepening understanding, curiosity, and enjoyment of Freemasonry. Mentoring does not end with the Third Degree; it is a lifelong opportunity to help Brethren appreciate Freemasonry - its purpose, traditions and meaning - rather than simply learn information."],
    ], [1800, 7612])
    b.key("These four principles are not sequential stages. They apply continuously and simultaneously throughout a member’s journey: from applicant to Entered Apprentice, through each degree, into the Royal Arch, into office, through the Chair and beyond it. What you Prepare an applicant for is entirely different from what you Prepare a Fellow Craft for; the principles stay constant while the content is always personal.", "Not stages")
    b.para("The Lodge Mentor co-ordinates; the Personal Mentor accompanies; but the responsibility belongs to every member of every Lodge. "
           "A Lodge that embeds these principles into its culture, not only into the work of its Lodge Mentor, is a Lodge in which members thrive and remain.", after=4)

    b.h1("3", "Who does what")
    b.para("Good mentoring depends on everyone knowing who does what. There are two distinct roles in the Lodge, a wider Provincial team behind them, and "
           "several other officers who each hold a piece of the work. Clarity prevents duplication and makes sure no member falls through the gaps.", after=6)
    b.h2("The member")
    b.para("At the centre. His needs for information, personal contact and one-to-one support drive everything. He sets the pace; he knows a Personal "
           "Mentoring Plan is kept, what it contains, and may see it on request; and he is never made to feel monitored, inadequate or incomplete.", after=6)
    b.h2("The Personal Mentor")
    b.para("One member, one-to-one. The member’s guide, friend and first point of contact through the early stages of his journey, and often beyond. "
           "He is not being asked to be an expert, a teacher or a ritualist. He is being asked to be a friend: the person a new Brother can ring, text or "
           "sit next to without feeling he is bothering anyone.", after=4)
    b.table([
        ["His role is", ["His first point of contact for any question", "Someone who keeps in touch between meetings", "A listener who notices when something changes", "His bridge to the Lodge and the Province"]],
        ["His role is not", ["The only person who can help him", "A ritual coach (that is the Lodge of Instruction)", "Someone who monitors or reports on him", "A replacement for the Lodge Mentor, who co-ordinates"]],
    ], [2200, 7212])
    b.para("Choose him for time, temperament and the potential for friendship, not for seniority. The Proposer often makes an excellent Personal Mentor "
           "because trust already exists; but not always, and the choice is the Lodge Mentor’s. Amwell Lodge’s practice of considered matching, on "
           "personality, shared interests, availability and the member’s particular needs, is the mature version of this.", after=4)
    b.key("The Personal Mentor’s most important quality is not knowledge. It is genuine interest in the wellbeing and progress of the person he is supporting. He is never required to know everything, never speculates, and is glad to say “I don’t know, but I’ll find out.”")
    b.h2("The Lodge Mentor")
    b.para("The co-ordinator. He makes sure mentoring happens, and happens well; he is not the Personal Mentor to every member. He:", after=3)
    b.bullets([
        "Ensures every new and joining member is assigned a suitable Personal Mentor, promptly, and gives each Personal Mentor his guide",
        "Opens a Personal Mentoring Plan for each member under mentoring and reviews progress against it",
        "Keeps an overview of mentoring across the Lodge and watches for patterns",
        "Is the Lodge’s link to the Province: the monthly call, his Assistant Provincial Grand Mentor, the quarterly picture",
        "Escalates concerns to the Worshipful Master or the Province when appropriate",
        "Keeps simple, secure records and briefs his successor",
        "Briefs the whole Lodge, at least once a year, on what mentoring is and why it matters",
    ])
    b.para("The Province recommends a sustained term, ideally three to five years. Continuity builds experience and avoids restarting the system with "
           "every change of appointment. Part of the job is to develop the man who will follow you.", after=4)
    b.h2("The Provincial Mentoring Team")
    b.para("The Provincial Grand Mentor, the Deputy Provincial Grand Mentor and five Assistant Provincial Grand Mentors, one for each APGM area. "
           "Your Assistant Provincial Grand Mentor is your nearest point of contact. He will visit, ideally at Initiations, and is there to help, never "
           "to inspect. Chapter 17 describes the support available.", after=6)
    b.h2("The other officers who hold a piece of it")
    b.table([
        ["Worshipful Master", "Sets the tone of the Lodge for a year and owns Step 3 conversations. Hears the Lodge Mentor’s report with the same weight as the Treasurer’s, briefs the Lodge on mentoring once a year, and protects junior Brethren from pressure in both directions."],
        ["Almoner", "Leads on welfare: health, money, bereavement and family, and on contact with non-attenders and former members. Welcomes new members jointly with the mentors. Anything touching welfare is referred to him; the mentoring record notes only “referred to Almoner”."],
        ["Membership Officer", "Holds the Lodge’s membership picture and the pipeline of candidates. Reviews joiners and leavers with the Lodge Mentor at least annually."],
        ["Director of Ceremonies", "Owns ceremony preparation and the culture of the Lodge of Instruction: how candidates are prepared practically, how juniors are drawn into ritual gently, and how the office queue is managed fairly."],
        ["Secretary", "Holds the candidate timeline that mentoring hangs from, and the attendance record, shared after every meeting. Gives same-day notice of any resignation, and acknowledges receipt only."],
        ["Treasurer", "Often the first to see hard evidence of disengagement. A first missed subscription earns a quiet word to the Lodge Mentor or Almoner, long before any formal process."],
        ["Royal Arch Representative", "One Journey’s bridge in the Lodge. Agrees with the Lodge Mentor what is said about the Royal Arch, when, and by whom."],
        ["Preceptor of the LoI", "Where not the DC: agrees how new members are welcomed at LoI and how learning is paced to the member rather than the calendar."],
    ], [2400, 7012])
    b.h2("The Mentor’s Jewel")
    b.para("In Hertfordshire the appointment of a Lodge Mentor is marked by the presentation and explanation of the Mentor’s Jewel: two chisels in saltire. "
           "The chisel is the tool of precision and of patient, deliberate work; the crossed chisels remind us that the mentor’s guidance and the member’s "
           "own effort are both necessary, and neither sufficient alone. The address ends with the sentence that belongs to every Brother in the room: "
           "we are all mentors to others in Freemasonry, and must take ownership of our future.", after=4)

    b.h1("4", "The mentor’s craft")
    b.para("Good mentoring is mostly good conversation. The content changes at every stage; the craft does not. This chapter is the Mentor Meeting Guide "
           "and the skills workshop in one place.", after=6)
    b.h2("Open questions, then silence")
    b.para("Begin with an open question and then listen. Do not fill the silence: count to three before speaking. He will usually fill it himself, and what "
           "he says then will be more honest than anything before it. Five questions that work:", after=3)
    b.bullets(["“What did you make of tonight?”", "“Is there anything from the ceremony you would like to talk through?”", "“How are you getting on with the Lodge of Instruction?”",
               "“Is there anything about Freemasonry you have been wondering about?”", "“And how are things generally? Everything alright?”"])
    b.para("The last one matters most. Ask it, mean it, and wait for the answer. When he raises a difficulty: listen without fixing, acknowledge it, ask what "
           "would help, and be honest when you do not know. Never guess on matters of Lodge governance or Masonic practice; find out, or point him to Solomon.", after=6)
    b.h2("The shape of a mentoring conversation")
    b.para("A framework, not a script. Structure matters because without it mentoring becomes reactive, addressing whatever comes up and neglecting what "
           "does not present itself. But structure ends where the person begins.", after=3)
    b.table([
        ["Before", "Where is he in his journey; what has happened since you last spoke; what is coming up for him and for the Lodge; anything outside the Lodge that might be affecting him."],
        ["Opening", "An open question, then silence. “What did you make of tonight?” “How have things been?” “What are you curious about at the moment?”"],
        ["Prepare", "Does he know what is coming next? If a ceremony is approaching: timing, format, general structure (never ritual detail), etiquette, dress, practicalities. Any concerns about time, cost or travel?"],
        ["Guide", "Has he been introduced to Solomon, the Provincial website, DAWG, the Fleet House Light Blues Club? Encouraged to visit, with company the first time? Does he understand how Lodge, Province and Grand Lodge fit together?"],
        ["Support", "The heart of it. Is he enjoying himself? Is anything not working? How is the balance with everything else? Does he feel welcomed and included? Watch for vagueness, deflection, or a drop in the enthusiasm of earlier conversations."],
        ["Enlighten", "What questions has he been carrying? “Is there anything from your Masonry recently that you’ve found yourself thinking about afterwards?” Follow his curiosity, not your own. Open doors; never stand in front of them; never over-explain."],
        ["Closing", "End with clarity: when you will next be in touch and anything either of you will do before then. Book the next contact before you part; never leave it open-ended."],
    ], [1700, 7712])
    b.h2("The relationship matures")
    b.para("In the early stages, from application to the Third Degree, regular contact matters most and the mentor is largely a teacher and a guide. As the "
           "member becomes established the formal frequency can reduce while the relationship deepens: teacher becomes coach, coach becomes advisor, advisor "
           "becomes the trusted friend he rings when something is wrong. The relationship never ends; its character simply evolves. And it runs both ways: "
           "the mentee teaches the mentor, and the mentor who stops learning from his Brother has stopped mentoring.", after=6)
    b.h2("Where your role ends")
    b.table([
        ["Welfare", "If he tells you something that touches his health, money, family or safety, speak to the Almoner and the Lodge Mentor promptly. You are not expected to carry it alone, and he will be better served if you do not."],
        ["Ritual", "Ritual is learned at the Lodge of Instruction, from the Preceptor and the DC. Encourage, accompany, reassure; do not coach what you have not been asked to coach."],
        ["Governance", "Never guess. Rules on resignation, attendance, dispensations and the like are in Chapter 18; anything else, ask the Secretary or the Provincial Grand Mentor’s office."],
        ["The Royal Arch", "Information, not recruitment. When he shows interest, bring in the Lodge Mentor and the Royal Arch Representative; never advocate for a particular Chapter’s need for members."],
    ], [1700, 7712])
    b.h2("Confidentiality and records")
    b.bullets([
        "Record only what mentoring needs: contact details, key dates, milestones and brief notes of support given. The Personal Mentoring Plan is the record; the Personal Mentor’s own notes are kept to a minimum and never shared beyond the Lodge Mentor",
        "The member knows a Plan is kept, what it contains, and may see it on request",
        "Welfare matters belong to the Almoner’s records. The mentoring Plan notes only “referred to Almoner”",
        "Keep records secure, paper or digital, and do not circulate them. Disclose confidential information only where the member’s interests require it or a safeguarding concern exists",
        "On handover, brief your successor and pass the Plans on; close and securely destroy a Plan within a defined period after a member resigns, migrates or dies",
    ])
    b.para("The Style, Terminology and Data Guide sets out the Province’s data rules for mentoring records in full.", italic=True, size=10, color=DARK, after=4)

    # ================================================================ PART TWO
    b.part("TWO", "The journey", "Stage by stage, from the first enquiry to the years beyond the Chair. Two rules run through every stage: there should never be any surprises, and the member sets the pace, never the calendar.")
    b.table([
        ["Stage", "Name", "Mentoring emphasis"],
        ["1", "Enquiry and application", "Honest expectation-setting; is the Lodge right for him; costs and time stated plainly."],
        ["2", "Before Initiation", "Personal Mentor assigned after ballot; practical preparation; contact the day before; introductions on the night."],
        ["3", "Entered Apprentice", "The critical window. Contact within a day or two; a one-to-one before the next meeting; the second meeting designed, not left to chance."],
        ["4", "Fellow Craft", "The forgotten degree given its due: knowledge, learning, readiness for the Third, at his pace."],
        ["5", "Master Mason", "The beginning, not the summit: visiting, participation, the wider structure, and the highest-risk drop-out point after the Third Degree."],
        ["6", "The Royal Arch: One Journey", "Completion, on his timing. Never a surprise, never a sales pitch. Information from Initiation onwards; the decision his alone."],
        ["7", "Into office", "Readiness, not pressure, in either direction. Ritual support without fear."],
        ["8", "The Chair", "Preparation for the year; support through it; and preparation, from Installation day, for what follows."],
        ["9", "Beyond the Chair", "Purpose after office: Included, Involved, Interested. Past Masters become Personal Mentors."],
        ["A", "At risk (cross-cutting)", "Warning signs, the four steps, the Rule 183 window, retrieval. Part Three."],
        ["B", "Relocation (cross-cutting)", "Migration as a continuation of mentoring. Part Three."],
    ], [800, 2600, 6012], header=True)

    b.h1("5", "Enquiry and application")
    b.para("Mentoring begins before the ballot, ideally from the moment a man first asks. The Lodge Mentor should be involved early, working with the "
           "Proposer, Seconder, Membership Officer and Secretary. A well-informed candidate makes a better long-term member; a surprised one is the first "
           "to leave.", after=6)
    b.h2("Before the interview")
    b.bullets([
        "Get to know him at informal social occasions before the formal interview, so that he arrives among friends rather than strangers",
        "Be open about what Freemasonry is and is not: a society of friends who meet, dine, support each other and give to charity; no secrets worth the name, and nothing he will be embarrassed to tell his family or his employer",
        "State the practical matters plainly: the joining fee and subscription, dining costs, meeting frequency and timing, the Lodge of Instruction, the time commitment if he later takes office, and the importance of a supportive partner or family",
        "Make sure he has introductory literature before the ballot: the Province’s Considering Freemasonry guide and the UGLE material",
        "Ask what draws him. Friendship, curiosity, service, tradition, a father or grandfather: the answer shapes everything the mentor does later",
    ])
    b.h2("The interview")
    b.para("Thorough and honest, in both directions. He is deciding about the Lodge as much as the Lodge is deciding about him. Cover what will be asked "
           "of him, in time and money, and what he can expect in return. Include his partner where he wishes; a partner who understands the commitment is "
           "the member’s strongest ally.", after=6)
    b.key("Is this the right Lodge for him? A man unhappy in one Lodge is often happy in another. If his circumstances, age or interests fit a neighbouring Lodge better, say so. Retrieval for the Craft matters more than recruitment to one Lodge.")

    b.h1("6", "Before Initiation, and the night itself")
    b.h2("After the ballot")
    b.bullets([
        "Assign a Personal Mentor as soon as possible after election, and open his Personal Mentoring Plan",
        "Meet him socially, with other members, as often as practical before his Initiation; involve his partner or family where appropriate",
        "Prepare him for the evening: timings, what he will be asked to do, what to wear, what to bring, and the general shape of the evening, without spoiling the ceremony itself",
        "The small human details: comfortable trousers, reading glasses in his pocket if he needs them, cash for the alms collection and raffle",
        "Any adjustments confirmed discreetly through the DC and Secretary: mobility, kneeling (not obligatory), faith requirements, the Volume of the Sacred Law appropriate to him",
        "A draft Festive Board reply provided well in advance, for him to adapt and make his own",
        "Contact him the day before, or on the morning: “We’re thinking of you. Arrive early. You’ll be guided at every step.”",
    ])
    b.h2("On the night")
    b.table([
        ["On arrival", ["Meet him at the door; he should never stand alone in the anteroom", "Introduce him to the Tyler and the Junior Deacon, who will look after him during the ceremony", "Introduce him to the Worshipful Master and DC if opportunity allows", "Reassure: everyone in the room has done this; the ceremony is symbolic, safe and dignified"]],
        ["During the meeting", ["Sit as close to him as custom allows once he is admitted; quiet explanation and reassurance as events unfold", "Someone is with him at every point he is outside the Lodge room"]],
        ["The Festive Board", ["Explain Masonic fire and the order of toasts before the meal; never let him struggle in front of the room", "Seat him according to Lodge custom and never among strangers", "Steady him before his reply; he has the draft"]],
        ["End of the evening", ["Congratulate him warmly and make sure others do", "Confirm the next Lodge date and your first one-to-one: book it now", "Present or signpost the Initiate’s Guide and the Province’s New Brother’s Guide, without overloading him"]],
    ], [2000, 7412])
    b.good("Introduce the candidate to the Tyler and the Junior Deacon immediately before the ceremony. Both can reassure him and make sure he does not feel isolated while he is outside the Lodge.")

    b.h1("7", "The Entered Apprentice")
    b.para("The evening of Initiation is one of the most formative experiences of a member’s Masonic life, and the weeks after it are the critical window. "
           "The role of Brother Initiate can feel very lonely: unfamiliar sights, unfamiliar language, unfamiliar people. Our task is to dispel that "
           "strangeness as quickly as possible, through warmth, explanation and genuine friendship.", after=6)
    b.h2("The first weeks")
    b.para("A quick call or message within a day or two, and a proper one-to-one before the next Lodge meeting. That first meeting is critical. Relaxed, "
           "at his pace, and including:", after=3)
    b.bullets([
        "An open question first: “What did you make of the evening?” Then let him talk",
        "A review of the ceremony: what happened and, in general terms, why. The Province’s First Degree Explained is written for exactly this conversation",
        "A chance to ask anything at all",
        "The test questions he will need before his Passing, and how he will be helped to learn them: at his pace, never the calendar’s",
        "How the Lodge is organised: officers, their roles, how to recognise them, and where the Lodge sits within Province and Grand Lodge",
        "Etiquette: the Risings, conduct at the Festive Board, the customs of the Lodge, dress (Chapter 19)",
        "The charitable work of Freemasonry, and how to sign up to Solomon",
    ])
    b.h2("His second meeting: the neglected danger point")
    b.para("A new Brother is treated as the hero of his Initiation, then too easily forgotten. His second meeting is often mostly business, and nobody "
           "has told him. He sits at dinner next to two men who talk across him about people he does not know. It is the moment most often left to "
           "chance, and the easiest to design.", after=3)
    b.table([
        ["Before", ["Contact him in the days before; confirm he is coming; answer the questions the first month has raised", "Decode the summons together: what each item means, how to reply, dining", "Regalia and dress checked"]],
        ["On the night", ["Meet him on arrival again; the anteroom is still unfamiliar", "Practise the salute and step quietly in a side room if he wishes; make sure he has the words of any odes", "Sit with him, or ensure a briefed Brother does; explain business items as they happen", "If the Lodge works a higher degree, an experienced Brother retires with him and uses the time well"]],
        ["The Festive Board", ["Seat him with company that will draw him in; do not let established groups close around him", "Introduce him to two or three members he has not properly met, by name and interest"]],
        ["Afterwards", ["Walk out with him; confirm the next contact and the next date", "Note anything that puzzled or unsettled him and cover it at your next one-to-one"]],
    ], [1900, 7512])
    b.h2("Between degrees")
    b.para("When the Lodge meets in a higher degree and he cannot be present, the evening is an opportunity for companionship and learning, not dead "
           "time. Meet him that evening or at another convenient time; use it for Masonic history, the meaning of the First Degree, the structure of "
           "Freemasonry, or whatever he is curious about; make sure he is exploring Solomon; and keep in frequent, informal contact. Regular check-ins "
           "are the difference between a member who feels included and one who quietly drifts.", after=6)
    b.h2("His first Lodge of Instruction")
    b.bullets([
        "Explain what it actually is: a relaxed rehearsal, officers in different seats, corrections, restarts, laughter. Unexpectedly ordinary, and that is its charm",
        "If there is little for an Entered Apprentice that week, say so honestly: “come anyway and I’ll introduce you to a few people”, or steer him to a better week",
        "Go with him, or arrange a named Brother to meet him at the door; introduce him to the Preceptor",
        "No pressure to take work. Attendance with the option not to perform is explicitly fine; involve him in something small if he is willing",
        "If he misses a week, the only right response on his return is warmth. LoI absence met with jokes or pointed remarks is a culture that loses members",
    ])
    b.h2("The summer, and the spring Initiate")
    b.para("National research finds that candidates initiated between January and June are less likely to remain than those initiated between September "
           "and December, largely because the summer break interrupts their early participation. Plan deliberate summer contact and a social touchpoint "
           "for every spring Initiate. The Lodge Mentor checks it has happened.", after=4)
    b.key("Isolation is one of the earliest predictors of disengagement. Belonging comes before knowledge: a man who feels he has friends in the room will learn everything else in time.")

    b.h1("8", "The Fellow Craft")
    b.para("The Second Degree is sometimes called the forgotten degree, a corridor between the First and the Third. It should not be. Its themes, "
           "knowledge, reason and the liberal arts and sciences, are the ones most easily connected to a member’s everyday life, and the Passing is the "
           "first time he answers in open Lodge.", after=6)
    b.h2("Before his Passing")
    b.bullets([
        "Test questions learned at his pace; rehearse together until he is comfortable, never rushed to fit the Lodge calendar",
        "Explain the tradition: he is tested on competence before advancement, and the questions are also a chance to discuss what the First Degree meant",
        "Timing, format and general structure of the ceremony, never ritual detail; reassure him that everyone finds answering in open Lodge daunting and rehearsal is the antidote",
        "Contact him the day before; meet him on arrival; re-introduce the Tyler, Inner Guard and Junior Deacon",
    ])
    b.h2("After his Passing")
    b.bullets([
        "Meet within days: “What did you make of it?” Let him talk first",
        "Talk through the differences he will have noticed between his Initiation and his Passing, and why they exist; the Province’s Second Degree Explained is the companion",
        "Discuss the liberal arts and sciences and their relevance to everyday life",
        "Make sure he can give the Sign, Token and Word of a Fellow Craft correctly, as taught in the Lodge; explain the different aprons",
        "Present and review the Fellow Craft’s Guide; encourage visiting, which he may now do freely",
        "Begin preparing him for the Third Degree, including the test questions, at his pace",
    ])
    b.good("Arrange for the Fellow Craft to witness a Passing in another Lodge as soon as convenient, and offer to go with him. Seeing the ceremony from the outside deepens understanding enormously.")
    b.key("Readiness, not convenience. If he is not ready to be passed, the right response is to wait and to fill the time well. What he will remember of a rushed ceremony is not the ceremony; it is the anxiety.")

    b.h1("9", "The Master Mason")
    b.para("The Raising is the most significant of the three degrees and a milestone, but it is a beginning, not a summit. The months after the Third "
           "Degree are one of the two recognised drop-out windows: the visible journey seems complete, the ceremonies have stopped, and nobody has shown "
           "him the road ahead. Mentoring intensifies here; it does not fade.", after=6)
    b.h2("Before his Raising")
    b.bullets([
        "Test questions rehearsed and comfortable",
        "Prepare him practically: timing, format, general structure; advise him to arrive well rested and having eaten",
        "Prepare him honestly for the ceremony’s weight: it is powerful, unlike what has come before, and entirely symbolic, safe and dignified",
    ])
    b.h2("After his Raising")
    b.bullets([
        "Check in within days; the reflective conversation matters more here than after any other degree. The central teaching, mortality, integrity and the life well lived, at his pace; the Province’s Third Degree Explained is the companion",
        "Make sure he can give the Grip and the Five Points of Fellowship correctly, as taught in the Lodge",
        "Visiting: how to sign in, the caution needed in Lodges under other Constitutions, respect for other Lodges’ workings (observe before judging; never criticise), the Visitors’ Toast and thanking hosts promptly",
        "The structure of Grand Lodge and Provincial Grand Lodge and their principal Rulers; the Grand Lodge Certificate, which is presented as soon as it arrives and is a passport, not for framing",
        "Present and review the Master Mason’s Guide",
    ])
    b.h2("The months after: show him the road ahead")
    b.bullets([
        "Lodge of Instruction, regularly and without obligation",
        "Visiting across the Province and beyond; go with him the first time",
        "Provincial events, the Fleet House Light Blues Club, DAWG, the Visitors’ Festival and the combined Mentoring and Learning and Development evenings",
        "Whatever in Freemasonry interests him: history, symbolism, the story of his own Lodge, charity, a piece of ritual he would like to learn",
        "The Royal Arch, explained as the completion of his journey, on his timing (Chapter 10)",
        "Gentle, open conversations about whether and when he might like to take office. Both answers are valid",
    ])
    b.key("Brethren who visit other Lodges, attend Provincial events and build friendships across the Province are more confident, better informed and far more likely to stay. Watch for any dip in enthusiasm now the ceremonies have finished; the Warning Signs card has the list.")

    b.h1("10", "One Journey: the Royal Arch")
    b.para("At the Union of 1813 Pure Antient Masonry was declared to consist of three degrees, including the Supreme Order of the Holy Royal Arch. UGLE "
           "describes the Royal Arch as the completion of that journey and encourages Master Masons to complete it before joining other orders. "
           "Hertfordshire’s approach is called One Journey. Its purpose is education and timing, never pressure. The most common failure is a Brother who "
           "reaches the Chair having never had the Royal Arch explained to him, and who then hears it as a sales pitch rather than the natural close of "
           "the path he began on his Initiation night.", after=6)
    b.h2("Four bridging points")
    b.table([
        ["At Initiation", "The Provincial Welcome Letter and the Personal Mentor’s early conversations frame Freemasonry as a four-part journey from the start. The Royal Arch is never a surprise."],
        ["Through progression", "As he is passed and raised, the mentor places each step in the wider journey, so the Royal Arch is anticipated rather than introduced cold."],
        ["Around the Third Degree", "The mentor makes sure he knows UGLE’s encouragement to complete the journey, and that the choice, and its timing, are his alone."],
        ["Through the Royal Arch Representative", "Every Craft Lodge has one. The Lodge Mentor works with him so that information is shared, accurate and consistent, and a Brother who expresses interest is introduced to a Chapter warmly."],
    ], [2600, 6812])
    b.h2("Two rules for the conversation")
    b.panel(["**Timing, not inevitability.** Introduce the Royal Arch when he has consolidated his understanding of the Craft, has shown curiosity or asked questions, and has room in his life for a further commitment. Never rush it, and never assume it. There is no universal right moment; there is only the right moment for him."], fill=PALE, bar=DEEP)
    b.panel(["**Information, not recruitment.** Tell him clearly what the Royal Arch is, what to expect, and what the time and cost involve. Reassure him that the Craft is complete and valuable in itself. He must never be made to feel an incomplete Freemason without it, and the moment a mentor advocates for a particular Chapter’s need for Companions, the mentoring relationship has been compromised."], fill=PALETEAL, bar=TEAL)
    b.h2("If he proceeds")
    b.para("Prepare him as carefully as for his Initiation: the format, location and general structure; the different officers and titles; dress and "
           "etiquette; the additional time and cost stated honestly. Support him through the joining process, and after Exaltation check in as carefully "
           "as after his Initiation. The Province’s Exaltation Explained is the companion. The Royal Arch is not the end of mentoring; it is simply the "
           "next conversation.", after=4)
    b.key("Members who join the Royal Arch are demonstrably more likely to stay engaged with their Craft Lodges. A Brother who sees a path ahead of him stays; one who believes he has reached the end of the road is the one most at risk.")

    b.h1("11", "Joining members, re-joiners and dual members")
    b.para("Not every new member is an Initiate. A Brother who joins from another Lodge arrives as an experienced Mason but a stranger to this one, and "
           "the risk is that the Lodge assumes he needs nothing. He does. He needs a friend, and he needs to learn how this Lodge works.", after=6)
    b.table([
        ["Joining members", ["A Personal Mentor for at least his first year, chosen as a guide to the Lodge rather than a teacher of Freemasonry", "A Personal Mentoring Plan marked as a joining-member route; the early-degree stages do not apply", "The things nobody thinks to explain: customs and working, dining, dress, who is who, how the Lodge communicates, the calendar", "Ask what he wants from his membership: office, friendship, a quieter Masonic life. Do not assume, and do not press him into the progression ladder", "If he has come from another Province, close the loop with the Provincial Grand Mentor’s office so the sending Lodge knows he has landed safely", "Introduce him at the Festive Board; he does not sit alone for his first few meetings"]],
        ["Re-joiners", ["A Brother returning after resignation is a success story in the making. Understand, gently, why he left and what is different now", "Treat him as a joining member for mentoring purposes, and be alert to whatever caused the earlier loss recurring", "Bring him up to date with changes in the Lodge, the Province and Grand Lodge practice, without fuss"]],
        ["Dual and multiple members", ["Watch the diary: two Lodges can each ask too much. Talk to his other Lodge Mentor where you can", "Never let two Lodges progress him to office at the same time unless he genuinely wants it", "Make sure he still feels he belongs here, and is not simply a visitor with a subscription"]],
    ], [2200, 7212])

    b.h1("12", "Into office, the Chair, and beyond it")
    b.h2("Into progressive office")
    b.para("When a member is offered his first office, confirm he is genuinely ready and willing; make sure he understands the responsibilities, time "
           "commitment and ritual requirements; arrange ritual support from experienced members and make sure he knows the DC and Preceptor are on his "
           "side; support regular attendance at the Lodge of Instruction; and check in as he progresses, alert to anxiety or fading confidence. "
           "Support must be available, not assumed.", after=4)
    b.table([
        ["Pushed too hard", ["Office before he feels ready or confident", "Ritual beyond his current ability", "Progression faster than he wishes", "The Royal Arch or other orders before the time is right", "Anything that makes him feel inadequate, unwelcome or judged"]],
        ["Left waiting", ["A capable Brother years from a first office behind a long queue, or behind Past Masters who prefer to hold the offices themselves", "A Brother never asked: to visit, to take a small piece of ritual, to help with a project", "A Past Master with nothing to do; a joining member whose experience nobody thought to use", "A Master Mason of several years who stopped learning because nobody offered him anything to learn"]],
    ], [2000, 7412])
    b.key("Readiness runs in both directions. A bored Brother leaves as surely as a frightened one. Talk to him about what he would like; find him a role that fits; raise the office queue with the Worshipful Master and the DC; and if the Lodge cannot offer what he needs, talk to the Province about a second Lodge. Outdated practices, sink or swim, discouraging visiting, unnecessary mystique, run counter to UGLE guidance; address them sensitively but directly.")
    b.h2("The Master Elect and the year in the Chair")
    b.para("The ritual book tells a Master what to say. Nothing tells him that the culture of his Lodge is, for one year, his; that the newest member’s "
           "whole impression of Freemasonry will be formed in that year; and that he is standing in both risk windows at once: his new members are in "
           "the first, and he is about to enter the second. Prepare him for the year: managing the Lodge, leading meetings, working with his officers, "
           "and looking after his own wellbeing. The Province’s session The Chair and Beyond does this for Wardens and Masters together.", after=4)
    b.table([
        ["1", "Set the tone", "What he laughs at, who he sits with, how he speaks to the newest member: the Lodge copies it. One quiet word about how returning members are greeted is worth a year of retention work."],
        ["2", "Hear the Lodge Mentor’s report", "On the Committee agenda, with the same weight as the Treasurer’s: who has a mentor and a Plan, who has been missed, what the Lodge could change."],
        ["3", "Brief the Lodge on mentoring", "Once in the year, in open Lodge. The Lodge Mentor has the presentation. We are all mentors to others."],
        ["4", "Protect juniors from pressure", "No office before ready; no ritual beyond ability; no Chapter sales pitch at the bar. And no capable man left waiting. The Master owns the office queue with the DC."],
        ["5", "Own the Step 3 conversation", "When a member is drifting and the matter touches Lodge culture, office or another member, the Master has the conversation, with the Almoner alongside for welfare."],
    ], [600, 2600, 6212])
    b.para("Family, Work, Masonry applies to the Master too. The year in the Chair is the point at which a Lodge is most likely to ask too much of one "
           "man. The Master who delegates, who lets the Wardens and the DC carry their share, and who protects his own family and work, finishes the year "
           "wanting more Freemasonry, not less. The one who does not is the one the Lodge loses in the year after.", after=6)
    b.h2("Beyond the Chair")
    b.para("Constitutionally, the Immediate Past Master is not an officer of the Lodge: the year after the Chair comes with a seat and no role, unless the "
           "Lodge gives him one. A seat is not a purpose. The Immediate Past Master’s year decides whether a Lodge keeps its Past Masters.", after=4)
    b.table([
        ["Included", "A seat, yes, and a voice: on the Committee, at the Lodge of Instruction, at the planning of the next year."],
        ["Involved", "A role or a project: the archive, the social calendar, a ceremony to work, a Brother to mentor. Agreed before Installation, not after."],
        ["Interested", "New questions: symbolism, history, the Royal Arch, the Province, a wider role. Welcome them, and prepare him for them."],
    ], [1800, 7612])
    b.key("Past Masters make the best Personal Mentors because they remember what the Chair cost and what it gave. Every Brother should want others to succeed, in Freemasonry and in life. The best Lodges are those in which experienced members invest in those coming behind them. That is mentoring in its truest sense.")

    from hb_more import more
    more(b)
    print("saved", b.save())


build()
