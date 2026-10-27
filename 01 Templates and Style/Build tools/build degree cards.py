"""One-page mentoring prompt cards per degree, on the Hertfordshire Mentoring template.
Usage: build_degree_cards.py <template.docx> <out_dir>"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from hmdocx import Builder, PRINCIPLES, DEEP, DARK, PALE, RGB
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

template, out_dir = sys.argv[1], sys.argv[2]

CARDS = {
    "Mentoring the Entered Apprentice": {
        "file": "Mentoring the Entered Apprentice One Page Issue 1 Draft.docx",
        "subtitle": "First Degree · prompts for the Personal Mentor",
        "intro": "Before, during and after his Initiation. Prompts, not a script. Explain the practical shape of the evening; never the ceremony itself.",
        "Prepare": [
            "Explain the timing, format and general structure of the evening, without giving away any ritual detail.",
            "Make sure he knows the etiquette and dress code: punctuality, mobile phone off, cash for the raffle and charity.",
            "Tell him he will be expected to say a few words at the Festive Board, and give him a short reply he can adapt.",
            "Practical reassurance: everyone has done this; arrive early, relax, you will be guided at every step.",
        ],
        "Guide": [
            "Introduce him to the Tyler, the Inner Guard and the Junior Deacon before the ceremony.",
            "Explain that Initiation is the symbolic beginning of his Masonic journey.",
            "Introduce the Masonic values of integrity, friendship, respect and service.",
            "Set out what is expected: attendance, contribution, charity and the Festive Board.",
            "Explain the social side and the importance of friendship, the Fleet House Light Blues Club, and signpost Solomon for further study.",
            "Encourage him to visit another Lodge’s Initiation.",
        ],
        "Support": [
            "Explain the journey that follows the First Degree, so there are no surprises.",
            "Questions are encouraged: you, the officers and the Brethren are available after the ceremony.",
            "Introduce the Lodge Mentor and explain what a mentor does, and does not do.",
            "Reassure him that the ceremony is symbolic, safe, and conducted with dignity and respect.",
            "Begin to identify the questions that will lead him towards his Second Degree.",
            "Explain what the Lodge of Instruction is and when it meets.",
        ],
        "Enlighten": [
            "Introduce the moral themes of the First Degree: Faith, Hope, Charity, Temperance, Fortitude, Prudence and Justice.",
            "Afterwards, encourage him to consider how the lessons of the First Degree apply to him personally.",
            "Emphasise the wider role of Freemasonry in supporting communities.",
            "Explain that there are three degrees, each building on the last, with opportunities for personal growth.",
            "Introduce the Daily Advancement WhatsApp Group (DAWG), run by the Provincial Mentoring Team.",
        ],
    },
    "Mentoring the Fellow Craft": {
        "file": "Mentoring the Fellow Craft One Page Issue 1 Draft.docx",
        "subtitle": "Second Degree · prompts for the Personal Mentor",
        "intro": "Before, during and after his Passing. Sometimes called the forgotten degree; it should not be. The pace is his, never the calendar’s.",
        "Prepare": [
            "Explain the timing, format and general structure of the evening, without giving away any ritual detail.",
            "Make sure he knows the etiquette and dress code: punctuality, mobile phone off, cash for the raffle and charity.",
            "Make sure he knows he will answer the questions in open Lodge, and offer help in practising them.",
            "Practical reassurance: everyone has done this; arrive early, relax, and trust that the officers will guide and support him throughout.",
        ],
        "Guide": [
            "Re-introduce him to the Tyler, the Inner Guard and the Junior Deacon.",
            "Explain that advancement builds on the symbolic foundation laid in the First Degree.",
            "Revisit the Masonic values of integrity, friendship, respect and service, and how they deepen through knowledge.",
            "Expectations: attendance, contribution, charity, and active participation at the Festive Board and the Lodge of Instruction.",
            "Encourage him to visit another Lodge’s Second Degree.",
        ],
        "Support": [
            "Explain that the Masonic journey continues beyond the Second Degree.",
            "Reinforce that questions are encouraged: you, the officers and the Brethren are available after the ceremony.",
            "Reassure him that ritual takes as long as each individual needs; he should not feel overwhelmed.",
            "Reassure him that the ceremony is symbolic, dignified and entirely safe.",
            "Begin to identify the areas of curiosity that will lead him towards his Third Degree.",
            "Highlight the value of the Lodge of Instruction.",
        ],
        "Enlighten": [
            "Discuss the moral and intellectual themes of the Second Degree: the pursuit of knowledge, the development of reason and the building of character.",
            "Encourage reflection on how the lessons of diligence, study and self-improvement apply personally and professionally.",
            "After the meeting, discuss the working tools of the first two degrees and their symbolic meaning.",
            "After the meeting, discuss the apron and what the Master Mason’s apron will signify.",
            "Keep him on the Daily Advancement WhatsApp Group (DAWG), run by the Provincial Mentoring Team.",
        ],
    },
    "Mentoring the Master Mason": {
        "file": "Mentoring the Master Mason One Page Issue 1 Draft.docx",
        "subtitle": "Third Degree · prompts for the Personal Mentor",
        "intro": "Before, during and after his Raising, and the months that follow. A beginning, not a summit, and a recognised drop-out window. Mentoring intensifies here.",
        "Prepare": [
            "Explain the timing, format and general structure of the evening, without giving away any ritual detail.",
            "Make sure he knows he will answer the questions in open Lodge, and offer help in practising them.",
            "Practical reassurance: the ceremony is solemn by design; he will be guided at every step, and everyone in the room has been where he is.",
        ],
        "Guide": [
            "Re-introduce him to the Tyler, the Inner Guard and the Junior Deacon.",
            "Explain how to sign in when visiting, and the care needed in Lodges under other Constitutions.",
            "Cover the etiquette of responding to the Visitors’ Toast, and of thanking hosts promptly.",
            "Explain that Lodges work differently: observe before judging, and never criticise another Lodge’s working.",
            "Check he understands the structure of Grand Lodge and Provincial Grand Lodge and knows their principal Rulers.",
            "Make sure he has received the Master Mason’s Guide.",
        ],
        "Support": [
            "Meet soon after the Raising. Review the evening and what it meant to him, at his pace.",
            "Encourage the Lodge of Instruction and visiting across the Province; go with him the first time.",
            "Introduce him to Provincial events, the Fleet House Light Blues Club and DAWG.",
            "Help him find the part of Freemasonry that interests him most, and support him in exploring it.",
            "Begin gentle, open conversations about whether and when he might like to take office. Both answers are valid.",
        ],
        "Enlighten": [
            "Discuss the essence of the Third Degree and its central teaching, with humility: Brethren read it differently, and none of them wrongly.",
            "Explain the Royal Arch as the completion of his journey in Pure Antient Masonry, on his timing: information, not recruitment. When he shows interest, bring in the Royal Arch Representative.",
            "Encourage personal learning through Solomon: history, symbolism, the story of his own Lodge.",
            "Remind him that mentoring does not end with the Third Degree, and that in time he will be a mentor to others.",
        ],
    },
}


def build(title, spec):
    out = os.path.join(out_dir, spec["file"])
    b = Builder(template, out, doc_title=title)
    doc = b.doc
    # compact title block
    tp = b.para(title, style="Title", after=2, before=0)
    for r in tp.runs: r.font.size = Pt(24)
    b.para(spec["subtitle"], bold=True, color="157A7C", size=10.5, after=3)
    b.para(spec["intro"], color=DARK, size=9.5, after=5)

    # a little more room on the page for the card: trim the bottom margin of this section
    sect = doc.sections[-1]
    from docx.shared import Twips
    sect.bottom_margin = Twips(1150)
    # four principle rows
    t = doc.add_table(rows=4, cols=2)
    b._widths(t, [1700, 7712])
    b._borders(t, {"insideH": ("single", 12, "FFFFFF"), "top": ("nil", 0, "auto"), "bottom": ("nil", 0, "auto")})
    b._margins(t, 70, 70, 150, 150)
    for i, (name, fill) in enumerate(PRINCIPLES):
        label = t.cell(i, 0); b._shade(label, fill)
        tcPr = label._tc.get_or_add_tcPr(); va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcPr.append(va)
        p = label.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        r = p.add_run(f"{i+1}"); r.font.color.rgb = RGB("8FD3D1"); r.bold = True; r.font.size = Pt(10)
        p2 = label.add_paragraph(); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.space_before = Pt(0)
        r = p2.add_run(name.upper()); r.font.color.rgb = RGB("FFFFFF"); r.bold = True; r.font.size = Pt(11)
        body = t.cell(i, 1); b._shade(body, PALE if i % 2 == 0 else "FFFFFF")
        b.cell_bullets(body, spec[name], after=0.5)
        for pp in body.paragraphs:
            pp.paragraph_format.line_spacing = 1.0
            for rr in pp.runs: rr.font.size = Pt(9.5)
    b._cant_split(t)

    b.para("Companion to the Personal Mentor’s Short Guide and the Stage Checklists · Never reveal ritual · "
           "Issue 1 (Draft for Review), October 2026 · © Provincial Grand Mentor, Province of Hertfordshire, 2026",
           size=8, color=DARK, before=5, after=0)
    print("saved", b.save())


for title, spec in CARDS.items():
    build(title, spec)
