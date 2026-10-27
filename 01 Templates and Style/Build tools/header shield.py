"""Place the shield in the first-page header, right-hand side of the identity row."""
import sys, copy
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out, shield = sys.argv[1], sys.argv[2], sys.argv[3]
doc = Document(src)
done = 0
for sec in doc.sections:
    hdr = sec.first_page_header
    for tbl in hdr.tables:
        for row in tbl.rows:
            if "Provincial Mentoring" not in row.cells[0].text and "Provincial Mentoring" not in "".join(c.text for c in row.cells):
                continue
            tcs = row._tr.findall(qn("w:tc"))
            if len(tcs) != 3:
                continue
            text_tc, right_margin = tcs[1], tcs[2]
            tcPr = text_tc.find(qn("w:tcPr"))
            span = tcPr.find(qn("w:gridSpan"))
            width = tcPr.find(qn("w:tcW"))
            SHIELD_W = 1730
            width.set(qn("w:w"), str(int(width.get(qn("w:w"))) - SHIELD_W))
            span.set(qn("w:val"), str(int(span.get(qn("w:val"))) - 1))
            # new cell for the shield
            new_tc = OxmlElement("w:tc")
            new_pr = copy.deepcopy(right_margin.find(qn("w:tcPr")))
            new_pr.find(qn("w:tcW")).set(qn("w:w"), str(SHIELD_W))
            for gs in new_pr.findall(qn("w:gridSpan")): new_pr.remove(gs)
            va = new_pr.find(qn("w:vAlign"))
            if va is None:
                va = OxmlElement("w:vAlign"); new_pr.append(va)
            va.set(qn("w:val"), "center")
            new_tc.append(new_pr)
            p = OxmlElement("w:p")
            pPr = OxmlElement("w:pPr")
            ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), "Header"); pPr.append(ps)
            jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "right"); pPr.append(jc)
            sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0"); pPr.append(sp)
            p.append(pPr)
            new_tc.append(p)
            right_margin.addprevious(new_tc)
            # add the picture through python-docx so the relationship is registered on the header part
            from docx.text.paragraph import Paragraph
            para = Paragraph(p, hdr._element)
            run = para.add_run()
            # run.add_picture needs the part: use the header's part via a temporary paragraph proxy
            from docx.parts.story import StoryPart
            part = hdr.part
            rId, image = part.get_or_add_image(shield)
            from docx.shared import Emu
            cx, cy = image.scaled_dimensions(None, Cm(2.9))
            inline = part.new_pic_inline(shield, None, Cm(2.9)) if hasattr(part, "new_pic_inline") else None
            if inline is None:
                from docx.oxml.shape import CT_Inline
                shape_id = part.next_id
                inline = CT_Inline.new_pic_inline(shape_id, rId, image.filename, cx, cy)
            run._r.add_drawing(inline)
            done += 1
doc.save(out); print("shield in header:", done, "->", out)
