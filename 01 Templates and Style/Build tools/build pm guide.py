"""Build B2a Personal Mentor's Short Guide on the Hertfordshire Mentoring Word template."""
import sys, copy, zipfile, shutil, re
from docx import Document
from docx.shared import Pt, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
TMP = OUT + ".stage1.docx"

# ---------- 1. fix headers/footers at XML level (canonical principle order) ----------
STRAP_OLD = "Prepare – Support – Guide – Enlighten"
STRAP_NEW = "Prepare • Guide • Support • Enlighten"
DOC_TITLE = "Personal Mentor’s Short Guide"

def patch(name, text):
    if name == "word/header1.xml" and STRAP_OLD in text:
        # strip labels are positional: swap the two middle labels, then fix the strapline
        text = text.replace(STRAP_OLD, "@@STRAP@@")
        text = text.replace(">Support<", ">@@G@@<").replace(">Guide<", ">Support<").replace(">@@G@@<", ">Guide<")
        text = text.replace("@@STRAP@@", STRAP_NEW)
    if name == "word/footer2.xml":
        text = text.replace(STRAP_OLD, STRAP_NEW)
        text = text.replace("Version [1.0]", "Issue 1 (Draft for Review)").replace("Review date [DD/MM/YYYY]", "October 2026")
    if name == "word/header2.xml":
        text = text.replace("[Document title or subject]", DOC_TITLE)
    return text

zin = zipfile.ZipFile(TEMPLATE)
zout = zipfile.ZipFile(TMP, "w", zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename.startswith("word/") and item.filename.endswith(".xml"):
        data = patch(item.filename, data.decode("utf-8")).encode("utf-8")
    zout.writestr(item, data)
zout.close(); zin.close()

# ---------- 2. body ----------
doc = Document(TMP)
body = doc.element.body
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# capture Layout 2 sectPr (the one with titlePg and footer rId7/rId8)
layout2 = None
for p in body.findall(qn("w:p")):
    sp = p.find(qn("w:pPr") + "/" + qn("w:sectPr"))
    if sp is not None and sp.find(qn("w:titlePg")) is not None:
        hdrs = [h.get(qn("r:id")) for h in sp.findall(qn("w:headerReference"))]
        ftrs = [f.get(qn("r:id")) for f in sp.findall(qn("w:footerReference"))]
        if "rId8" in ftrs:
            layout2 = copy.deepcopy(sp)
assert layout2 is not None
for child in list(body):
    body.remove(child)
for f in layout2.findall(qn("w:footerReference")):
    if f.get(qn("w:type")) == "first": f.set(qn("r:id"), "rId7")
body.append(layout2)

INK, DEEP, DARK, TEAL, PALE, PALETEAL = "1E2329", "003668", "06345D", "157A7C", "F2F7FC", "E6F2F2"
PRINCIPLES = [("Prepare", "003668"), ("Guide", "0A4F72"), ("Support", "106576"), ("Enlighten", "157A7C")]

def para(text="", style=None, bold=False, italic=False, color=None, size=None, after=None, before=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        add_runs(p, text, bold=bold, italic=italic, color=color, size=size)
    pf = p.paragraph_format
    if after is not None: pf.space_after = Pt(after)
    if before is not None: pf.space_before = Pt(before)
    if align: p.alignment = align
    if keep: pf.keep_with_next = True
    return p

def add_runs(p, text, bold=False, italic=False, color=None, size=None):
    """**bold** inline markup supported."""
    parts = re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text)
    for part in parts:
        if not part: continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]; b = True
        elif part.startswith("*") and part.endswith("*"):
            part = part[1:-1]; i = True
        r = p.add_run(part)
        r.bold = b or None
        r.italic = i or None
        if color: r.font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor.from_string(color)
        if size: r.font.size = Pt(size)
    return p

def h1(num, text, fill=None):
    p = doc.add_paragraph(style="Heading1")
    if num:
        r = p.add_run(f" {num} ")
        rpr = r._r.get_or_add_rPr()
        c = OxmlElement("w:color"); c.set(qn("w:val"), "FFFFFF"); rpr.append(c)
        s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill or DEEP); rpr.append(s)
        p.add_run("  " + text)
    else:
        p.add_run(text)
    return p

def h2(text):
    return doc.add_paragraph(text, style="Heading2")

def bullet(text):
    p = doc.add_paragraph(style="ListBullet"); add_runs(p, text); return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)

def cell_margins(tbl, top=100, bottom=100, left=170, right=170):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tblPr.append(m)

def borders(tbl, spec):
    """spec: dict side -> (val, sz, color)"""
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        val, sz, col = spec.get(side, ("nil", 0, "auto"))
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), col); b.append(e)
    tblPr.append(b)

def set_widths(tbl, widths):
    tbl.autofit = False
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(w))
    for row in tbl.rows:
        for c, w in zip(row.cells, widths): c.width = Twips(w)
    tblW = tbl._tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tbl._tbl.tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tbl._tbl.tblPr.append(lay)

def look(tbl, first_row=0, first_col=1):
    tblPr = tbl._tbl.tblPr
    tl = tblPr.find(qn("w:tblLook"))
    if tl is None:
        tl = OxmlElement("w:tblLook"); tblPr.append(tl)
    val = (0x0020 if first_row else 0) | (0x0080 if first_col else 0) | 0x0400 | 0x0200
    tl.set(qn("w:val"), "%04X" % val)
    tl.set(qn("w:firstRow"), str(first_row)); tl.set(qn("w:lastRow"), "0")
    tl.set(qn("w:firstColumn"), str(first_col)); tl.set(qn("w:lastColumn"), "0")
    tl.set(qn("w:noHBand"), "1"); tl.set(qn("w:noVBand"), "1")

def cant_split(tbl):
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)

def cell_text(cell, text, bold=False, color=None, size=None, first=True):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
    add_runs(p, text, bold=bold, color=color, size=size)
    return p

def cell_bullets(cell, items):
    first = True
    for it in items:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.style = doc.styles["ListBullet"]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(0)
        add_runs(p, it)

def panel(text_lines, fill=PALE, bar=TEAL, after=6):
    """Boxed key-principle panel: single cell, pale fill, thick left bar."""
    t = doc.add_table(rows=1, cols=1)
    set_widths(t, [9412])
    borders(t, {"left": ("single", 24, bar)})
    cell_margins(t, 90, 90, 200, 200)
    c = t.cell(0, 0); shade(c, fill)
    first = True
    for line in text_lines:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
        add_runs(p, line, color=DARK)
    cant_split(t)
    para(after=after - 6, size=4)
    return t

def principles_strip():
    t = doc.add_table(rows=2, cols=4)
    set_widths(t, [2353, 2353, 2353, 2353])
    borders(t, {"insideV": ("single", 12, "FFFFFF")})
    cell_margins(t, 90, 90, 140, 140)
    defs = {
        "Prepare": "Ready for each stage before it happens. No surprises.",
        "Guide": "Accurate, honest, timely information. Never dependency.",
        "Support": "Present, responsive, compassionate. “How can we help?”",
        "Enlighten": "Understanding, curiosity and enjoyment, for life.",
    }
    for i, (name, fill) in enumerate(PRINCIPLES):
        c = t.cell(0, i); shade(c, fill)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        r = p.add_run(f"{i+1}  "); r.font.color.rgb = RGB("8FD3D1"); r.bold = True; r.font.size = Pt(11)
        r = p.add_run(name.upper()); r.font.color.rgb = RGB("FFFFFF"); r.bold = True; r.font.size = Pt(11)
        c2 = t.cell(1, i); shade(c2, PALE)
        p2 = c2.paragraphs[0]; p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.space_before = Pt(0)
        r = p2.add_run(defs[name]); r.font.size = Pt(9.5); r.font.color.rgb = RGB(DARK)
    cant_split(t)
    return t

from docx.shared import RGBColor
def RGB(h): return RGBColor.from_string(h)

# ---------- content ----------
para("Practitioner guide · B2a", style="DocumentType")
para("Personal Mentor’s Short Guide", style="Title")
para("For the Brother who has just been asked to look after a new member: what the role is, what to do at each stage, "
     "and where to turn. The full detail is in the Personal Mentor’s Guide; this is the version you can read on the train.",
     style="PurposeStatement")

meta = doc.add_table(rows=2, cols=2, style="Mentoring Table")
set_widths(meta, [2041, 7371])
for i, (k, v) in enumerate([
    ("Audience", "Personal Mentors, newly appointed or returning. Useful to Lodge Mentors when briefing them."),
    ("Issue", "Issue 1 (Draft for Review), October 2026 · Deputy Provincial Grand Mentor, on behalf of the Provincial Grand Mentor"),
]):
    cell_text(meta.cell(i, 0), k); cell_text(meta.cell(i, 1), v)
look(meta, 0, 1); cant_split(meta)
para(after=0, size=4)

# 1
h1("1", "What you are being asked to do", DEEP)
para("You are not being asked to be an expert, a teacher or a ritualist. You are being asked to be a friend: "
     "the person a new Brother can ring, message or sit next to without feeling he is bothering anyone. "
     "Everything else in this guide is that friendship, applied at the right moments.", after=6)

role = doc.add_table(rows=2, cols=2, style="Mentoring Table")
set_widths(role, [4706, 4706])
cell_text(role.cell(0, 0), "Your role is", bold=True, color="FFFFFF")
cell_text(role.cell(0, 1), "Your role is not", bold=True, color="FFFFFF")
shade(role.cell(0, 0), DEEP); shade(role.cell(0, 1), DEEP)
_cb = cell_bullets
def cell_bullets_tight(cell, items):
    _cb(cell, items)
    for pp in cell.paragraphs: pp.paragraph_format.space_after = Pt(1)
cell_bullets_tight(role.cell(1, 0), [
    "His first point of contact for any question",
    "Someone who keeps in touch between meetings",
    "A listener who notices when something changes",
    "His bridge to the Lodge and the Province",
])
cell_bullets_tight(role.cell(1, 1), [
    "The only person who can help him",
    "A ritual coach. That is the Lodge of Instruction",
    "Someone who monitors or reports on him",
    "A replacement for the Lodge Mentor, who co-ordinates",
])
look(role, 1, 0); cant_split(role)
for c in role.rows[0].cells:
    for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
para(after=0, size=4)

para("**The four principles.** Mentoring in Hertfordshire rests on four principles, defined once in "
     "*The Hertfordshire Mentoring Principles*. They run together, all the time. They are not stages.", after=6, keep=True)
principles_strip()
para(after=0, size=4)

# 2
h1("2", "Getting started", DEEP)
para("The Lodge Mentor will introduce you. Before you meet, find out who proposed him, what he does, whether there is "
     "a partner or family to include, and anything practical such as distance, shift work or travel. "
     "A word with his Proposer is worth ten minutes of guessing.", after=8)
para("**Your first meeting sets the tone.** Coffee, a pint or a chat after a meeting is fine. Aim to:", after=4, keep=True)
for b in [
    "Introduce yourself properly, including what surprised or confused you when you started.",
    "Explain simply what you are for: a point of contact, a guide, a friend.",
    "Ask what drew him to Freemasonry and what he hopes to get from it. Then listen.",
    "Say the four things that matter: there are no stupid questions; call me any time; this is meant to be enjoyable, so tell me if it is not; and nobody will push you before you are ready.",
    "Agree how and when you will next be in touch.",
]: bullet(b)

# 3
h1("3", "The journey at a glance", DEEP)
para("Prompts, not a script. Use what you know of the man to decide what he needs and when. "
     "Full detail for every stage is in the Personal Mentor’s Guide, and one-page checklists exist for the key nights.", after=8)

journey = [
    ("Before Initiation", [
        "Meet him socially; include his partner or family where appropriate.",
        "Make sure he knows what to wear, what to bring, the timings and what the evening involves in practical terms.",
        "Call or message him the day before.",
    ]),
    ("Initiation night", [
        "Introduce him to the Tyler and the Junior Deacon before the meeting.",
        "Sit as close to him as you can. Be the face he recognises.",
        "Prepare him for the Festive Board: the toasts, Masonic Fire, and a short reply he has seen in advance.",
        "Never let him stand alone.",
    ]),
    ("First weeks as an Entered Apprentice", [
        "A quick call or message within a day or two: “What did you make of it?”",
        "A proper one-to-one before the next Lodge meeting. Review the evening at his pace, not yours.",
        "Over the first few meetings cover who is who in the Lodge, etiquette and the Festive Board, charity, and the Initiate’s Guide.",
        "Introduce Solomon and the Lodge of Instruction.",
        "Discuss visiting, the Fleet House Light Blues Club and the Daily Advancement WhatsApp Group (DAWG).",
    ]),
    ("Between degrees", [
        "When the Lodge works in a higher degree he cannot attend. This is the lonely stretch. Keep in touch so he never feels forgotten.",
        "Use the time for learning: history, the Lodge room, his questions.",
        "Keep an eye on his progress towards Passing, at his pace, never the calendar’s.",
    ]),
    ("After Passing", [
        "Meet soon afterwards; talk through the ceremony and what it meant to him.",
        "Hand over the Fellow Craft’s Guide and discuss it.",
        "Encourage him to see a Passing elsewhere; begin gentle preparation for the Third.",
    ]),
    ("After Raising", [
        "The beginning, not the summit, and the highest-risk drop-out point. Meet soon.",
        "Cover visiting etiquette, signing in, replying to the Visitors’ Toast, and how Lodge, Province and Grand Lodge fit together.",
        "Hand over the Master Mason’s Guide. Go visiting together.",
    ]),
    ("Life as a Master Mason", [
        "Help him find the part of Freemasonry that interests him most.",
        "The Royal Arch: information, not recruitment; timing. When he shows interest, bring in the Royal Arch Representative.",
    ]),
    ("Into office and beyond", [
        "Readiness, not pressure, in either direction. Never let him accept an office before he feels able.",
        "Watch for anxiety about ritual and offer practical help.",
        "Stay in touch after the Chair. Past Masters drift too.",
    ]),
]
jt = doc.add_table(rows=len(journey), cols=2, style="Mentoring Table")
set_widths(jt, [2300, 7112])
for i, (stage, items) in enumerate(journey):
    cell_text(jt.cell(i, 0), stage)
    cell_bullets(jt.cell(i, 1), items)
look(jt, 0, 1); cant_split(jt)
para(after=0, size=4)

# 4
h1("4", "Conversations that matter", DEEP)
para("Good mentoring is mostly good conversation. Open questions, then silence. Five that work:", after=4, keep=True)
for q in [
    "“How did you find last night?”",
    "“Is there anything from the ceremony you would like to talk through?”",
    "“How are you getting on with the Lodge of Instruction?”",
    "“Is there anything about Freemasonry you have been wondering about?”",
    "“And how are things generally? Everything alright?”",
]: bullet(q)
para("The last one matters most. Ask it, mean it, and wait for the answer. When he raises a difficulty: listen without "
     "fixing, acknowledge it, ask what would help, and be honest when you do not know. Never guess on matters of "
     "Lodge governance or Masonic practice.", before=6, after=8)
panel(["**Welfare.** If he tells you something that touches his health, money, family or safety, speak to the Lodge "
       "Almoner and the Lodge Mentor promptly. You are not expected to carry it alone, and he will be better served if you do not."],
      fill=PALETEAL, bar=TEAL)

# 5
h1("5", "Noticing when something changes", DEEP)
para("Two missed meetings without a word. Going quiet on messages. Dropping out of the Lodge of Instruction. Repeated "
     "talk of time or money pressure. A flat mood at the Festive Board. One of these is nothing. A pattern is a conversation.", after=8)
para("The first step is always warm and informal: a call, a message, a coffee. If it continues, tell the Lodge Mentor "
     "that week. Early action keeps members; late action rarely does. The Warning Signs card has the full list.", after=8)
panel(["**The right question.** Always “How can I help?”, never “Why haven’t you been attending?”. "
       "The tone decides whether he feels supported or judged."], fill=PALE, bar=DEEP)

# 6
h1("6", "Help for you", DEEP)
help_rows = [
    ("Lodge Mentor", "Your first call for anything you are unsure about. He co-ordinates mentoring across the Lodge and can escalate if needed."),
    ("Lodge Almoner", "Any welfare concern: health, finances, bereavement, family. Involve him early."),
    ("Provincial Grand Mentor’s office", "For anything that needs Provincial support or advice, in confidence. mentor@hertfordshirefreemasons.org.uk"),
    ("Monthly mentoring call", "First Tuesday of the month, 7pm, on Teams. Open to every Lodge Mentor and Personal Mentor. Bring questions."),
    ("Solomon", "UGLE’s learning platform and the place to send him for accurate answers: solomon.ugle.org.uk"),
]
ht = doc.add_table(rows=len(help_rows), cols=2, style="Mentoring Table")
set_widths(ht, [2600, 6812])
for i, (k, v) in enumerate(help_rows):
    cell_text(ht.cell(i, 0), k); cell_text(ht.cell(i, 1), v)
look(ht, 0, 1); cant_split(ht)
para(after=0, size=4)
para("Keep a simple, private note of when you spoke and what you covered. It helps you, and it helps whoever follows you. "
     "Do not share it beyond the Lodge Mentor. If in doubt about whether something is serious, say so to the Lodge Mentor. "
     "Raising a small thing costs nothing; missing a big one costs a Brother.", after=10)
panel(["**Family, Work, Masonry: in that order, always.**",
       "Freemasonry must be fun. A man who is enjoying himself attends; a man who attends participates; "
       "a man who participates belongs; and a man who belongs does not resign. That is the whole job."],
      fill=PALE, bar=TEAL)

para("Provincial Office, Fleet House, 10 Parkway, St Albans AL3 6PA · 01727 833061 · www.pgherts.org · solomon.ugle.org.uk  |  "
     "Companion to: Personal Mentor’s Guide (Issue 1a, July 2026) · Mentor Meeting Guide · Stage Checklists · Warning Signs card  |  "
     "© Provincial Grand Mentor, Province of Hertfordshire, 2026", size=8.5, color=DARK, before=4, after=0)

doc.save(OUT)
import os; os.remove(TMP)
print("saved", OUT)
